# -*- coding: utf-8 -*-
"""Gera o documento (.docx) do estudo com janelas do 2º IMDC.
   Data-driven: lê resultados/*.json e imdc_ref/metrics_wis.csv. Roda após o runner concluir."""
import os, sys, json
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

AQUI = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(AQUI, "resultados"); FIG = os.path.join(AQUI, "figuras")
OUT = os.path.join(AQUI, "Estudo_IMDC_janelas.docx")
DESK = r"C:/Users/julia/OneDrive/Área de Trabalho/2026_07_17-Estudo_IMDC_janelas.docx"

ORDEM = [("Clássicos", ["RF", "XGBoost", "LSTM", "SARIMA", "EnsRFXGB"]),
         ("Quânticos", ["QRC", "QSVM", "VQR_ReUp", "VQR_Angle", "QLSTM"]),
         ("Ensembles", ["EnsRegime"])]
ROTULO = {"RF": "Random Forest", "XGBoost": "XGBoost", "LSTM": "LSTM clássico", "SARIMA": "SARIMA",
          "EnsRFXGB": "Ensemble RF+XGBoost", "QRC": "QRC", "QSVM": "QSVM", "VQR_ReUp": "VQR Re-uploading",
          "VQR_Angle": "VQR AngleEmbedding", "QLSTM": "QLSTM", "EnsRegime": "Ensemble por regime (QSVM→RF)"}


def carregar():
    d = {}
    for f in os.listdir(RES):
        if f.endswith("_resultados.json"):
            d[f.replace("_resultados.json", "")] = json.load(open(os.path.join(RES, f), encoding="utf-8"))
    return d


def imdc_ref():
    w = pd.read_csv(os.path.join(AQUI, "imdc_ref", "metrics_wis.csv"))
    wdf = w[w.state == "DF"].copy(); wdf["validation_test"] = wdf["validation_test"].astype(str)
    out = {}
    for vt, t in [("1", "T1"), ("2", "T2"), ("3", "T3")]:
        s = wdf[wdf.validation_test == vt]["WIS"]
        out[t] = {"melhor": float(s.min()), "mediana": float(s.median()), "n": int(s.count())}
    return out


def cell(c, txt, bold=False, fill=None, size=9, align="left"):
    c.text = ""; p = c.paragraphs[0]; p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    r = p.add_run(str(txt)); r.bold = bold; r.font.size = Pt(size)
    if fill:
        from docx.oxml.ns import qn; from docx.oxml import OxmlElement
        sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), fill); c._tc.get_or_add_tcPr().append(sh)


def tabela_metricas(doc, dados):
    modelos = [m for _, ms in ORDEM for m in ms if m in dados]
    t = doc.add_table(rows=1 + len(modelos), cols=8); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["Modelo", "T1 R²", "T1 WIS", "T2 R²", "T2 WIS", "T3 R²", "T3 WIS", "Tipo"]
    for j, h in enumerate(hdr): cell(t.rows[0].cells[j], h, bold=True, fill="1F3864", size=9, align="center")
    for j in range(8): t.rows[0].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tipo = {m: g for g, ms in ORDEM for m in ms}
    for i, m in enumerate(modelos, 1):
        d = dados[m]; row = t.rows[i].cells
        def g(c, k): return d.get(c, {}).get(k)
        vals = [ROTULO.get(m, m)]
        for c in ["T1", "T2", "T3"]:
            r2 = g(c, "R2"); wis = g(c, "WIS")
            vals += [f"{r2:.3f}" if r2 is not None else "—", f"{wis:.0f}" if wis is not None else "—"]
        vals.append(tipo[m])
        for j, v in enumerate(vals): cell(row[j], v, bold=(j == 0), align="left" if j in (0, 7) else "center")
    return t


def tabela_comparacao(doc, dados, ref):
    linhas = [m for m in ["RF", "XGBoost", "EnsRFXGB", "QRC"] if m in dados]
    t = doc.add_table(rows=3 + len(linhas), cols=4); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(["Modelo / Referência", "WIS T1", "WIS T2 (surto)", "WIS T3"]):
        cell(t.rows[0].cells[j], h, bold=True, fill="1F3864", size=9, align="center")
        t.rows[0].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, m in enumerate(linhas, 1):
        d = dados[m]; cell(t.rows[i].cells[0], ROTULO.get(m, m), bold=True)
        for j, c in enumerate(["T1", "T2", "T3"], 1):
            cell(t.rows[i].cells[j], f"{d[c]['WIS']:.0f}", align="center")
    r = len(linhas)
    cell(t.rows[r + 1].cells[0], "2º IMDC — melhor equipe (DF)", bold=True, fill="FBE0D2")
    cell(t.rows[r + 2].cells[0], "2º IMDC — mediana (DF)", bold=True, fill="FBF3CE")
    for j, c in enumerate(["T1", "T2", "T3"], 1):
        cell(t.rows[r + 1].cells[j], f"{ref[c]['melhor']:.0f}", align="center", fill="FBE0D2")
        cell(t.rows[r + 2].cells[j], f"{ref[c]['mediana']:.0f}", align="center", fill="FBF3CE")


def main():
    dados = carregar(); ref = imdc_ref()
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2); sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Replicação do estudo com as janelas do 2º IMDC", level=0)
    p = doc.add_paragraph(); r = p.add_run(
        "Benchmark clássico × quântico de previsão de dengue no Distrito Federal reproduzido sobre as janelas "
        "de validação do 2º Infodengue-Mosqlimate Dengue Challenge (Testes 1, 2 e 3), permitindo comparação "
        "direta com as 17 equipes que submeteram modelos para o DF. Alvo: casos_est. Métrica primária: WIS.")
    r.italic = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_heading("1. Janelas de validação", level=1)
    doc.add_paragraph("T1 (temporada típica): 2022-10-09 → 2023-10-01 · 52 semanas · treino 36 sem.\n"
                      "T2 (surto recorde 2024): 2023-10-08 → 2024-09-29 · 52 semanas · treino 88 sem.\n"
                      "T3 (pós-surto): 2024-10-06 → 2025-09-28 · 52 semanas · treino 140 sem.\n"
                      "Treino expansivo desde jan/2022; cada teste é uma temporada completa.")

    doc.add_heading("2. Desempenho de todos os modelos", level=1)
    tabela_metricas(doc, dados)

    doc.add_heading("3. Comparação com o placar do 2º IMDC (DF)", level=1)
    tabela_comparacao(doc, dados, ref)
    cav = doc.add_paragraph()
    cr = cav.add_run("Ressalva metodológica: este estudo prevê um passo à frente usando os lags reais de cada "
                     "semana, ao passo que as equipes do IMDC preveem a temporada inteira a partir de um único "
                     "ponto de origem. Isso favorece este estudo nos regimes estáveis (T1, T3) e reduz a vantagem "
                     "no surto (T2), onde a magnitude do pico escapa a ambos.")
    cr.italic = True; cr.font.size = Pt(9); cr.font.color.rgb = RGBColor(0x80, 0x40, 0x00)

    doc.add_heading("4. Figuras — Predição vs. Observado", level=1)
    for grupo, ms in ORDEM:
        ms = [m for m in ms if m in dados]
        if not ms: continue
        doc.add_heading(grupo, level=2)
        for m in ms:
            fp = os.path.join(FIG, f"{m}_pred_vs_obs.png")
            if os.path.exists(fp):
                hp = doc.add_paragraph(); hr = hp.add_run(ROTULO.get(m, m)); hr.bold = True; hr.font.size = Pt(12)
                ip = doc.add_paragraph(); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ip.add_run().add_picture(fp, width=Cm(15.5))
                doc.add_paragraph()

    doc.save(OUT)
    try:
        import shutil; shutil.copy(OUT, DESK); print("copiado p/ Área de Trabalho:", DESK)
    except Exception as e:
        print("aviso copia:", e)
    print("SALVO:", OUT, "| modelos:", len(dados))


if __name__ == "__main__":
    main()
